import os
import random
import docx
from lxml import etree
import csv
import openpyxl


def parse_file(file_path):
    """
    Parses the given file and returns the extracted data.
    """
    file_ext = os.path.splitext(file_path)[1].lower()

    try:
        if file_ext == ".csv":
            return parse_csv(file_path)
        elif file_ext in [".xlsx", ".xls"]:
            return parse_excel(file_path)
        elif file_ext == ".txt":
            return parse_txt(file_path)
        elif file_ext == ".xml":
            return parse_xml(file_path)
        elif file_ext == ".docx":
            return parse_docx(file_path)
        else:
            return ["Unsupported file type"]
    except Exception as e:
        return [f"Error parsing file: {e}"]


def parse_csv(file_path):
    try:
        with open(file_path, 'r', newline='', encoding='utf-8') as csvfile:
            reader = csv.DictReader(csvfile)
            rows = list(reader)
        sample_size = max(1, int(0.1 * len(rows)))
        return random.sample(rows, sample_size)
    except Exception as e:
        return [f"Error parsing CSV: {e}"]


def parse_excel(file_path):
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        headers = [str(cell.value) if cell.value else f"Column{i}" for i, cell in enumerate(sheet[1])]
        rows = []

        for row in sheet.iter_rows(min_row=2):
            row_data = {}
            for i, cell in enumerate(row):
                value = cell.value if cell.value is not None else ""
                key = headers[i] if i < len(headers) else f"Column{i}"
                row_data[key] = value
            rows.append(row_data)

        sample_size = max(1, int(0.1 * len(rows)))
        return random.sample(rows, sample_size)
    except Exception as e:
        return [f"Error parsing Excel: {e}"]


def parse_txt(file_path):
    try:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                lines = f.readlines()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='latin-1') as f:
                lines = f.readlines()

        sample_size = max(1, int(0.1 * len(lines)))
        sample_lines = random.sample(lines, sample_size)
        return [{"Line": line.strip()} for line in sample_lines]
    except Exception as e:
        return [f"Error parsing TXT: {e}"]


def parse_xml(file_path):
    try:
        tree = etree.parse(file_path)
        root = tree.getroot()
        text_list = [element.text for element in root.iter() if element.text]
        lines = "\n".join(text_list).splitlines()

        sample_size = max(1, int(0.1 * len(lines)))
        return random.sample(lines, sample_size)
    except Exception as e:
        return [f"Error parsing XML: {e}"]


def parse_docx(file_path):
    try:
        doc = docx.Document(file_path)
        lines = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        sample_size = max(1, int(0.1 * len(lines)))
        return random.sample(lines, sample_size)
    except Exception as e:
        return [f"Error parsing DOCX: {e}"]
